from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parents[1] / "Nomad_Wanderers_Demo_Document_With_Diagrams.docx"


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def add_numbered(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Number")


def add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Light Shading Accent 1"
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value


def add_diagram(document: Document, diagram: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.style = "Normal"
    run = paragraph.add_run(diagram)
    run.font.name = "Consolas"
    run.font.size = Pt(9)


def main() -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    for style_name in ("Title", "Heading 1", "Heading 2"):
        styles[style_name].font.name = "Aptos Display"
        styles[style_name].font.color.rgb = RGBColor(0, 92, 94)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Nomad Wanderers\nApplication Demo Document")
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(0, 92, 94)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Tour discovery, booking workflow, admin management and demo payment flow").italic = True
    document.add_paragraph("Version: Demo build | Prepared for product demonstration", style="Subtitle")

    document.add_heading("1. Demo Objective", level=1)
    document.add_paragraph(
        "Nomad Wanderers is a travel-experience application for discovering and booking locally led tours in Indian cities. "
        "The demo shows the public traveller journey and the secure admin workflow for managing tours, enquiries, custom journeys and demo payments."
    )

    document.add_heading("2. Product Overview", level=1)
    add_bullets(document, [
        "Public website for browsing curated local experiences.",
        "Trips planning by city and trip duration: One-day trips and Weekly trips.",
        "Tour booking flow with customer details and a safe demo payment checkout.",
        "Design your journey form for tailor-made itineraries.",
        "Admin dashboard for managing operational travel data.",
        "Responsive navigation and presentation across desktop and mobile views.",
    ])

    document.add_heading("3. Public Traveller Experience", level=1)
    add_table(document, ["Area", "What the traveller can do"], [
        ["Home", "View the full-width tour carousel, featured travel stories, and clear calls to action."],
        ["Tours", "Search experiences and filter by interest such as Heritage, Food, Culture, Walking, Adventure, Shared or Private."],
        ["Trips", "Choose One-day or Weekly trips, select Mumbai or Delhi, and view the matching published tours."],
        ["Booking", "Choose a tour, provide contact details, continue to the demo checkout, and receive a demo confirmation screen."],
        ["Design your journey", "Submit destinations, dates, duration, travellers, budget and interests for a custom itinerary."],
        ["Contact", "Send a general support, partnership or travel question to the Nomad team."],
    ])

    document.add_heading("4. Tour and Trip Concept", level=1)
    document.add_paragraph("Tours and Trips have different purposes in the application:")
    add_bullets(document, [
        "A Tour is a specific, bookable experience, for example: Mumbai Heritage Walk.",
        "Trips are a discovery layer that organizes tours by city and duration.",
        "A tour is created once in admin and appears under Trips automatically according to City and Trip type.",
        "Tour format is Shared or Private. Trip type is One-day trip or Weekly trip.",
    ])

    document.add_heading("5. Admin Dashboard Capabilities", level=1)
    add_table(document, ["Admin area", "Current demo capability"], [
        ["Tours", "Create, edit, publish, unpublish, feature and delete tours. Manage image, city, trip type, format, category, duration, price and highlights."],
        ["Contact enquiries", "Review general support messages, update statuses and add private follow-up notes."],
        ["Custom journeys", "Review Design your journey submissions, add itinerary notes, quotations and progress status."],
        ["Demo payments", "Review completed simulated booking payments, payment method, amount, traveller and transaction reference."],
        ["Search and pagination", "Search every admin data area and view six records per page using server-side API pagination."],
    ])

    document.add_heading("6. Booking and Demo Payment Flow", level=1)
    add_numbered(document, [
        "Traveller selects Book this tour from a tour card, carousel or trip page.",
        "Traveller enters name, email, phone/WhatsApp number, preferred date and traveller count.",
        "The application redirects to a Razorpay-style demo checkout page.",
        "Traveller selects UPI, Credit card or Debit card as the simulated payment method.",
        "Traveller clicks the demo payment button for a ₹499 booking-deposit simulation.",
        "The backend creates a demo payment record with a unique DEMO transaction reference.",
        "The traveller sees email and WhatsApp confirmation previews on the success screen.",
        "The admin sees the record in the Demo payments tab.",
    ])
    document.add_paragraph(
        "Important demo note: this flow does not collect real money and does not send live email or WhatsApp messages. "
        "It is intentionally labelled as a simulation. A production flow must use Razorpay payment verification through webhooks."
    )

    document.add_heading("6.1 High-Level Architecture Flow", level=2)
    add_diagram(document, """TRAVELLER / ADMIN
        |
        v
React + Vite Frontend
Home | Tours | Trips | Booking | Admin Dashboard
        |
        | HTTPS API requests
        v
FastAPI Backend (main.py)
Public APIs | JWT-protected Admin APIs | Pagination | Validation
        |
        | Data-access layer
        v
MySQL Database (db_models.py)
Tours | Enquiries | Custom Journeys | Demo Payments | Revoked Tokens
        |
        +--> Demo checkout confirmation preview (current demo)
        |
        +--> Razorpay + Email + WhatsApp services (future production integration)""")

    document.add_heading("6.2 End-to-End Booking Sequence Diagram", level=2)
    add_diagram(document, """Traveller        React Frontend       FastAPI Backend         MySQL            Admin Dashboard
    |                    |                    |                   |                    |
    | Browse tours/trips |                    |                   |                    |
    |------------------->| GET /api/tours     |                   |                    |
    |                    |------------------->| SELECT tours      |                    |
    |                    |                    |------------------>|                    |
    |                    |<-------------------| paginated results |                    |
    |<-------------------| tour cards          |                   |                    |
    |                    |                    |                   |                    |
    | Book this tour     |                    |                   |                    |
    |------------------->| booking details    |                   |                    |
    |                    | redirect to demo checkout                |                    |
    | Choose UPI/card    |                    |                   |                    |
    |------------------->| POST /api/demo-payments                  |                    |
    |                    |------------------->| INSERT payment     |                    |
    |                    |                    |------------------>|                    |
    |                    |<-------------------| demo reference     |                    |
    |<-------------------| confirmation preview|                   |                    |
    |                    |                    |                   |                    |
    |                    |                    | GET admin payments |                    |
    |                    |                    |<------------------|                    |
    |                    |                    |----------------------------------->|
    |                    |                    |      payment record visible         |
    |                    |                    |                   |                    |""")

    document.add_heading("7. Search and Pagination", level=1)
    add_table(document, ["Page", "Records per page", "Filtering method"], [
        ["Admin dashboard", "6", "Server-side search and pagination for Tours, Enquiries, Custom journeys and Demo payments."],
        ["Public Tours", "9", "Server-side search plus experience-category and Shared/Private filters."],
        ["Public Trips", "6", "Server-side city and trip-type filters with pagination."],
    ])
    document.add_paragraph("Example API pattern:", style="Heading 2")
    document.add_paragraph("/api/admin/tours?page=1&page_size=6&search=mumbai")
    document.add_paragraph("/api/tours?city=Mumbai&trip_type=One-day+trip&page=1&page_size=6")

    document.add_heading("8. Backend and Data Design", level=1)
    add_bullets(document, [
        "FastAPI backend with routes exposed from main.py.",
        "Database access and schema functions are maintained in db_models.py.",
        "MySQL connection details are read on the server from backend/db.env.",
        "JWT-based admin authentication protects all admin management APIs.",
        "Core tables: tours, contact_enquiries, custom_journeys, demo_payments and revoked_tokens.",
        "Public APIs accept only the data required for browsing, enquiries, custom journeys and demo payments.",
    ])

    document.add_heading("9. Suggested Live Demo Walkthrough", level=1)
    add_numbered(document, [
        "Open the Home page and show the tour carousel and responsive navigation.",
        "Open Tours and demonstrate experience search and category/format filters.",
        "Open Trips, choose One-day trips, select Mumbai, and show paginated results.",
        "Select Book this tour and enter traveller details.",
        "Complete the UPI/Card dummy payment and show the confirmation reference.",
        "Sign in to Admin Dashboard.",
        "Open Demo payments to show the saved simulated payment record.",
        "Open Enquiries and Custom journeys to demonstrate status, notes and quotation management.",
        "Open Tours in admin and create or edit a tour with City, Trip type and Tour format.",
        "Use admin search and pagination to demonstrate handling larger data volumes.",
    ])

    document.add_heading("10. Production Readiness: Next Steps", level=1)
    add_bullets(document, [
        "Correct the MySQL username/password and permissions in backend/db.env so the backend can start and create tables.",
        "Replace static demo admin credentials with user accounts, hashed passwords and role-based access control.",
        "Integrate Razorpay Key ID, Key Secret and webhook secret for real payment collection and verified payment status.",
        "Integrate an email provider for confirmation emails and a WhatsApp Business provider/template for verified notifications.",
        "Add dedicated customer profiles, guide scheduling, promotions, content management and reports as the next admin modules.",
        "Add audit logs, backups, production environment variables and monitoring before launch.",
    ])

    document.add_heading("11. Demo Summary", level=1)
    document.add_paragraph(
        "Nomad Wanderers demonstrates a complete travel platform foundation: visitors discover meaningful local experiences, "
        "browse trips by destination and duration, submit tailored journey requests, complete a safe simulated booking payment, "
        "and the admin team manages the operational data from one dashboard."
    )

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
