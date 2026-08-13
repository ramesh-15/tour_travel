from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parents[1] / "Nomad_Wanderers_Two_Phase_Project_Scope.docx"


def bullets(document: Document, values: list[str]) -> None:
    for value in values:
        document.add_paragraph(value, style="List Bullet")


def table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    result = document.add_table(rows=1, cols=len(headers))
    result.style = "Light Shading Accent 1"
    for index, value in enumerate(headers):
        result.rows[0].cells[index].text = value
    for row in rows:
        cells = result.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value


def main() -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    document.styles["Normal"].font.name = "Aptos"
    document.styles["Normal"].font.size = Pt(10.5)
    for style_name in ("Title", "Heading 1", "Heading 2"):
        document.styles[style_name].font.name = "Aptos Display"
        document.styles[style_name].font.color.rgb = RGBColor(0, 92, 94)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Nomad Wanderers\nTwo-Phase Travel Booking Platform Scope")
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(0, 92, 94)
    document.add_paragraph("Travel experiences, tour discovery, booking management and payment integration", style="Subtitle")

    document.add_heading("Project Overview", level=1)
    document.add_paragraph(
        "Nomad Wanderers is a travel-experience booking platform designed for curated local tours, one-day experiences, "
        "weekly trips and tailor-made journeys. The platform gives travellers a refined discovery and booking experience "
        "while providing the Nomad team with one secure administration dashboard."
    )
    document.add_paragraph(
        "The project is planned in two phases. Phase 1 delivers the core traveller and admin platform. Phase 2 completes "
        "the production payment, confirmation and operational integration layer."
    )

    document.add_heading("PHASE 1 — Core Tour Discovery and Booking Platform", level=1)
    document.add_paragraph("Goal: Launch a polished, responsive booking platform for tours and custom travel enquiries.")

    document.add_heading("Traveller-Facing Features", level=2)
    bullets(document, [
        "Responsive home page with a full-width featured-tour carousel, travel storytelling sections and clear booking calls to action.",
        "Desktop and mobile navigation with Tours, Trips, About, Contact and Design your journey access.",
        "Tours catalogue with search and filters for Heritage, Food, Culture, Walking, Adventure, Shared and Private experiences.",
        "Trips discovery pages for One-day trips and Weekly trips, organized by city and trip type.",
        "Server-side pagination for fast browsing of larger tour collections.",
        "Book this tour journey with traveller name, email, phone/WhatsApp number, preferred date and traveller count.",
        "Design your journey form for destinations, travel dates, duration, travellers, budget and special interests.",
        "Contact form for general support, travel questions and partnerships.",
        "Social-media footer links and WhatsApp contact access.",
    ])

    document.add_heading("Admin Panel Features", level=2)
    table(document, ["Admin area", "Management capability"], [
        ["Tours", "Add, edit, publish, unpublish, feature and delete tours. Manage images, city, trip type, Shared/Private format, category, duration, price, highlights and badges."],
        ["Bookings", "View booking requests, review traveller details, update booking status and manage confirmation workflow."],
        ["Custom journeys", "Review tailor-made requests, prepare itineraries, add quotations and follow-up notes."],
        ["Contact enquiries", "Read general messages, update statuses and add internal notes."],
        ["Search and pagination", "Search and paginate Tours, Bookings, Enquiries and Custom journeys for efficient daily operations."],
        ["Content controls", "Manage featured tour selection and carousel-ready tour content from the tour management workflow."],
    ])

    document.add_heading("Phase 1 Technical Foundation", level=2)
    bullets(document, [
        "React/Vite responsive frontend.",
        "FastAPI backend with JWT-protected admin APIs.",
        "MySQL database with separate schema/data-access code and API routes.",
        "Server-side validation, search and pagination for public and admin collections.",
        "Data model for tours, traveller enquiries, custom journeys, bookings, payments and admin session security.",
    ])

    document.add_heading("PHASE 2 — Live Payments, Confirmations and Operational Integrations", level=1)
    document.add_paragraph("Goal: Convert the booking platform into a production-ready payment and confirmation system.")

    document.add_heading("Razorpay Payment Integration", level=2)
    bullets(document, [
        "Secure Razorpay Checkout integration for UPI, credit cards, debit cards, net banking and supported wallets.",
        "Full-payment or configurable partial-advance/deposit payment option for every tour.",
        "Server-generated Razorpay orders so pricing and booking data remain secure.",
        "Razorpay webhook verification before a payment is marked as successful.",
        "Booking status updates: Payment pending, Partially paid, Paid, Failed, Refunded or Cancelled.",
        "Automatic receipt/invoice reference linked to each successful payment.",
        "Admin payment view for deposits, balance due, refunds and transaction references.",
    ])

    document.add_heading("Automatic Confirmation Notifications", level=2)
    bullets(document, [
        "Booking confirmation email sent after the Razorpay payment webhook is verified.",
        "WhatsApp Business template message sent to the traveller's provided WhatsApp number after confirmation.",
        "Confirmation content includes traveller name, booked tour, date, traveller count, payment status, receipt reference and support contact details.",
        "Admin notification for every new paid booking so the operations team can assign guides and prepare the experience.",
        "Optional reminder notifications before the tour date and balance-payment reminders where applicable.",
    ])

    document.add_heading("Future Operations Extension", level=2)
    bullets(document, [
        "Customer profiles with booking history, preferences, notes and repeat-traveller details.",
        "Guide profiles, languages, availability, trip assignment and capacity management.",
        "Promotions, discount codes, featured seasonal campaigns and announcement banners.",
        "Content management for About content, FAQ, testimonials, contact details, WhatsApp number and social links.",
        "Operational reports for booking totals, popular tours/cities, enquiry sources and conversion rates.",
    ])

    document.add_heading("Production Flow", level=1)
    document.add_paragraph("Traveller selects a tour → enters booking details → pays securely through Razorpay → Razorpay webhook verifies the payment → booking is marked paid → receipt is created → confirmation email and WhatsApp template message are sent → admin receives the paid booking for operations follow-up.")

    document.add_heading("Required Production Credentials", level=1)
    bullets(document, [
        "MySQL database hostname, database name, username and password.",
        "Razorpay Key ID, Key Secret and webhook secret.",
        "Email service credentials and approved sender email address.",
        "WhatsApp Business API credentials, phone number ID, access token and approved template name/language.",
        "Final policy for full payment, percentage deposit, cancellation and refund rules.",
    ])

    document.add_heading("Project Outcome", level=1)
    document.add_paragraph(
        "At completion, Nomad Wanderers will have a scalable travel booking platform where travellers can discover meaningful "
        "experiences, book and pay securely, receive verified confirmations, and where the operations team can manage tours, "
        "bookings, enquiries, custom journeys and payments from one professional admin dashboard."
    )

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
